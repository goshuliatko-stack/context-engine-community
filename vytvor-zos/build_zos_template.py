# -*- coding: utf-8 -*-
"""
ŠABLÓNA pre generovanie Zmluvy o sprostredkovaní (ZoS) cez python-docx.

POUŽITIE:
  0. PRED úpravou tohto skriptu prejdi povinný rozhovor s používateľom (SKILL.md sekcia 0):
     sprostredkovateľ, cena, provízia, doba platnosti, počet predávajúcich, manželia/BSM,
     body 9.–12. čl. III, typ nehnuteľnosti (byt/dom/pozemok) — a podľa odpovedí nastav
     DATA["je_bsm"], DATA["vratane_ponuky_9_12"], DATA["nehnutelnost_typ"] a SPROSTREDKOVATEL nižšie.
  1. pip install python-docx
  2. Uprav blok DATA nižšie (osoby, nehnuteľnosť, cena, provízia, doba).
  3. Neznáme údaje nechaj ako R("[DOPLNIŤ]") — vykreslia sa ČERVENO.
  4. python3 build_zos_template.py  →  vygeneruje .docx
  5. Pošli súbor cez SendUserFile a skús upload do cieľovej Drive zložky.

POZNÁMKY:
  - Právny text čerpaj z VZORU (read_file_content), nie z hlavy.
  - Pri VIAC vlastníkoch sklonuj "Záujemca" → "Záujemcovia/Záujemcov/Záujemcom/Záujemcami".
  - Sprostredkovateľ NIE je automaticky Venoc s. r. o. — vždy sa najprv spýtaj (pozri SKILL.md sekcia 0).
  - Ak su predávajúci manželia → DATA["je_bsm"] = True → podiel 1/1 (BSM), nie zlomky.
"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

# ============================================================================
# DATA — UPRAV TENTO BLOK PRE KONKRÉTNY OBCHOD
# ============================================================================
DATA = {
    "cena_cislom": "149 990,- EUR",
    "cena_slovom": "jednostoštyridsaťdeväťtisícdeväťstodeväťdesiat EUR",
    "provizia_pct": "5,23 %",
    "doba_dni": "150",
    "miesto_podpisu": "Žiline",
    "pocet_vyhotoveni": "3",   # = počet zmluvných strán
    "out_path": "/tmp/contract_work/Zmluva_o_sprostredkovani.docx",
    # Sú predávajúci manželia? Ak True → BSM (bezpodielové spoluvlastníctvo manželov), podiel 1/1.
    # Ak False → bežné podielové spoluvlastníctvo (podiel každého z poľa "podiel" nižšie).
    "je_bsm": False,
    # Majú byť v zmluve body 9.–12. čl. III ("predám tvoju nehnuteľnosť alebo dovolenka zdarma")?
    "vratane_ponuky_9_12": True,
    # Typ nehnuteľnosti: "byt" | "dom" | "pozemok" — ovplyvňuje formulácie nižšie.
    "nehnutelnost_typ": "byt",
}

# Záujemcovia (spoluvlastníci). Pridaj/uber položky podľa počtu vlastníkov.
# Pole s hodnotou None → vykreslí sa ČERVENO ako [DOPLNIŤ].
ZAUJEMCOVIA = [
    {
        "meno": "Tomáš Schmidt",
        "rodne_priezvisko": None,             # ak rovnaké, daj None / vynechaj
        "datum_narodenia": "03.11.1994",
        "rodne_cislo": None,                  # [DOPLNIŤ]
        "trvaly_pobyt": "Hlinská 2589/20, 010 01 Žilina, SR",
        "rodinny_stav": None,                 # [DOPLNIŤ]
        "mail": "tomas.schmidt013@gmail.com",
        "tel": "+421 904 531 346",
        "ucet": None,                         # [DOPLNIŤ – IBAN]
        "podiel": "1/2",
    },
    {
        "meno": "Natália Schmidtová",
        "rodne_priezvisko": "Petráková",
        "datum_narodenia": "20.09.1991",
        "rodne_cislo": None,
        "trvaly_pobyt": "Komenského 1310/30, 024 01 Kysucké Nové Mesto, SR",
        "rodinny_stav": "vydatá",             # potvrdené používateľom → čierno
        "mail": "petrakovanatalia35@gmail.com",
        "tel": "+421 911 477 114",
        "ucet": None,
        "podiel": "1/2",
    },
]

# Sprostredkovateľ — VŽDY si najprv potvrď s používateľom, či má byť Venoc s. r. o. alebo iná firma
# (pozri SKILL.md sekcia 0, otázka č. 1). Údaje Venoc s. r. o. sú uložené aj v Context Engine
# (ctx_company("Venoc s. r. o.")) — skús ich odtiaľ dohľadať ako prvý zdroj.
SPROSTREDKOVATEL = {
    "obchodne_meno": "Venoc s. r. o.",
    "or_zapis": "zapísaná v Obchodnom registri Okresného súdu Žilina, oddiel: Sro, vl. č. 79531/L",
    "ico": "54 555 680",
    "sidlo": "Gabajova 2593/22, 010 01 Žilina",
    "ucet": "SK56 8330 0000 0026 0217 8024",
    "konatel": "Georgi Conev",
    "mail": "conevreality@gmail.com",
    "tel": "+421 948 880 069",
    "web": "www.georgiconev.sk",
}

# Popis nehnuteľnosti z LV — uprav podľa výpisov listu vlastníctva.
NEHNUTELNOST = (
    "bytu číslo 2, nachádzajúceho sa na prízemí, vo vchode číslo 1, stavby bytový dom, so súpisným "
    "číslom 3447, zapísaného na liste vlastníctva číslo 1738, pre katastrálne územie Závodie, obec "
    "Žilina, okres Žilina, okresný úrad Žilina, katastrálny odbor (ďalej ako aj „byt“) postaveného na "
    "pozemku KN C parcelné číslo 1472/51, druh pozemku Zastavaná plocha a nádvorie, o výmere 629 m², "
    "zapísanom na liste vlastníctva číslo 2924, pre katastrálne územie Závodie, obec Žilina, okres "
    "Žilina, okresný úrad Žilina, katastrálny odbor (ďalej ako aj „pozemok“),"
)
PODIEL_SPOLOCNE = "4984/357640"
TITUL_NADOBUDNUTIA = "Kúpnej zmluvy, vklad povolený dňa 6.9.2021 pod č. V-7849/2021 (774/2021)"
TARCHY = (
    "záložné právo v prospech Slovenská sporiteľňa, a. s., Tomášikova 48, 832 37 Bratislava, "
    "IČO 00151653, pod č. V-6894/2021 a záložné právo v prospech Slovenská sporiteľňa, a. s., pod "
    "č. V-11329/2022. "
)

# ============================================================================
# INFRAŠTRUKTÚRA — väčšinou netreba meniť
# ============================================================================
RED = RGBColor(0xFF, 0x00, 0x00)
BLACK = RGBColor(0x00, 0x00, 0x00)
FONT = "Times New Roman"

doc = Document()
style = doc.styles["Normal"]
style.font.name = FONT
style.font.size = Pt(12)
rPr = style.element.get_or_add_rPr()
rFonts = rPr.find(qn('w:rFonts'))
if rFonts is None:
    rFonts = rPr.makeelement(qn('w:rFonts'), {})
    rPr.append(rFonts)
rFonts.set(qn('w:eastAsia'), FONT)

for s in doc.sections:
    s.left_margin = Cm(2.5)
    s.right_margin = Cm(2.5)
    s.top_margin = Cm(2)
    s.bottom_margin = Cm(2)


def add_p(segments, align=None, bold=False, italic=False, size=12, space_after=8, space_before=0):
    """segments: string, (text, red_bool) alebo list (text, red_bool)."""
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if isinstance(segments, str):
        segments = [(segments, False)]
    elif isinstance(segments, tuple):
        segments = [segments]
    for text, red in segments:
        run = p.add_run(text)
        run.font.name = FONT
        run.font.size = Pt(size)
        run.bold = bold
        run.italic = italic
        run.font.color.rgb = RED if red else BLACK
    return p


def heading(text, size=14):
    return add_p(text, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=size, space_after=4, space_before=10)


def body(segments, justify=True, space_after=8, size=12):
    align = WD_ALIGN_PARAGRAPH.JUSTIFY if justify else None
    return add_p(segments, align=align, space_after=space_after, size=size)


R = lambda t: (t, True)    # červený placeholder
N = lambda t: (t, False)   # normálny text


def field(label, value, placeholder="[DOPLNIŤ]"):
    """Pole 'Label: hodnota' — ak value je None, hodnota je červená."""
    if value is None:
        body([N(label + ": "), R(placeholder)], justify=False, space_after=2)
    else:
        body(N(f"{label}: {value}"), justify=False, space_after=2)


# Označenie nehnuteľnosti v texte podľa typu (byt/dom/pozemok) — pozri DATA["nehnutelnost_typ"]
NEHNUTELNOST_LABEL = {
    "byt": "byt a pozemok",
    "dom": "dom a pozemok",
    "pozemok": "pozemok",
}[DATA["nehnutelnost_typ"]]

# Pomocníci na skloňovanie (jeden vs viac vlastníkov)
PLURAL = len(ZAUJEMCOVIA) > 1
Z_NOM = "Záujemcovia" if PLURAL else "Záujemca"          # 1. pád
Z_GEN = "Záujemcov" if PLURAL else "Záujemcu"            # 2. pád
Z_DAT = "Záujemcom" if PLURAL else "Záujemcovi"          # 3. pád
Z_INS = "Záujemcami" if PLURAL else "Záujemcom"          # 7. pád
SU = "sú" if PLURAL else "je"
MAJU = "majú" if PLURAL else "má"
VYHLASUJU = "vyhlasujú" if PLURAL else "vyhlasuje"

# ============================================================================
# TELO ZMLUVY
# ============================================================================
add_p("Zmluva o sprostredkovaní", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=18, space_after=6)
add_p("uzavretá v zmysle ust. § 774 a nasl. Zák. č. 40/1964 Zb. Občianskeho zákonníka v platnom znení nasledovne:",
      align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, size=11, space_after=16)

# --- Článok I. Zmluvné strany ---
heading("Článok I.", 14)
heading("Zmluvné strany", 12)

for i, z in enumerate(ZAUJEMCOVIA, start=1):
    field("Meno, Priezvisko", z["meno"])
    if z.get("rodne_priezvisko"):
        field("Rodné priezvisko", z["rodne_priezvisko"])
    field("Dátum narodenia", z.get("datum_narodenia"))
    field("Rodné číslo", z.get("rodne_cislo"))
    field("Trvalý pobyt", z.get("trvaly_pobyt"))
    body(N("Štátna príslušnosť: Slovenská republika"), justify=False, space_after=2)
    field("Rodinný stav", z.get("rodinny_stav"))
    field("Mail", z.get("mail"))
    field("Tel. číslo", z.get("tel"))
    field("Č. účtu", z.get("ucet"), placeholder="[DOPLNIŤ – IBAN]")
    podiel_zobrazeny = "1/1 (bezpodielové spoluvlastníctvo manželov – BSM)" if DATA["je_bsm"] else z.get("podiel")
    field("Spoluvlastnícky podiel na nehnuteľnosti", podiel_zobrazeny)
    body(N(f"Ďalej len ako „Záujemca {i}“."), justify=False, space_after=10)

if PLURAL:
    body(N("(Záujemca 1 a Záujemca 2 spolu ďalej aj ako „Záujemcovia“ v príslušnom gramatickom tvare.)"),
         justify=False, space_after=10)
    if DATA["je_bsm"]:
        body(N(f"{Z_NOM} {VYHLASUJU}, že sú manželmi a nehnuteľnosť nadobudli a vlastnia v bezpodielovom "
               "spoluvlastníctve manželov (ďalej len „BSM“) v podiele 1/1."), justify=False, space_after=10)

s = SPROSTREDKOVATEL
body(N(f"Obchodné meno: {s['obchodne_meno']}"), justify=False, space_after=2)
body(N(s["or_zapis"]), justify=False, space_after=2)
body(N(f"IČO: {s['ico']}"), justify=False, space_after=2)
body(N(f"Sídlo: {s['sidlo']}"), justify=False, space_after=2)
body(N(f"Č. účtu: {s['ucet']}"), justify=False, space_after=2)
body(N(f"Konateľ: {s['konatel']}"), justify=False, space_after=2)
body(N(f"Mail: {s['mail']}"), justify=False, space_after=2)
body(N(f"Tel. číslo: {s['tel']}"), justify=False, space_after=2)
body(N("Ďalej len ako „Sprostredkovateľ“."), justify=False, space_after=10)

body(N(f"(spolu {Z_NOM} a Sprostredkovateľ ďalej aj ako „zmluvné strany“)"), space_after=10)
body(N("ktorí vyhlásili, že sú k uzatváranému právnemu úkonu oprávnení a k právnym úkonom v plnom rozsahu spôsobilí, a to za týchto podmienok:"), space_after=14)

# --- Článok II. Úvodné ustanovenia ---
heading("Článok II.", 14)
heading("Úvodné ustanovenia", 12)
body(N("Sprostredkovateľ je právnická osoba vykonávajúca podnikateľskú činnosť v zmysle príslušných oprávnení uvedených v Obchodnom registri SR."))
if DATA["je_bsm"]:
    podiel_veta = f"{Z_NOM} {SU} manželmi a vlastníkmi v bezpodielovom spoluvlastníctve manželov (BSM) v podiele 1/1 nasledujúcej nehnuteľnosti:"
elif PLURAL:
    podiel_veta = f"{Z_NOM} {SU} podielovými spoluvlastníkmi, každý v rozsahu spoluvlastníckeho podielu, nasledujúcej nehnuteľnosti:"
else:
    podiel_veta = f"{Z_NOM} {SU} vlastníkom nasledujúcej nehnuteľnosti:"
body(N(podiel_veta))
body(N(NEHNUTELNOST))
body(N(f"Podiel priestoru na spoločných častiach a spoločných zariadeniach domu, na príslušenstve a spoluvlastnícky podiel k pozemku: {PODIEL_SPOLOCNE}."))
body(N(f"(ďalej spoločne {NEHNUTELNOST_LABEL} aj ako „nehnuteľnosť“)"))
body(N(f"{Z_NOM} {VYHLASUJU}, že nehnuteľnosť nadobudli na základe {TITUL_NADOBUDNUTIA}."))
body([N(f"{Z_NOM} {VYHLASUJU}, že na nehnuteľnosti viaznu nasledovné ťarchy: {TARCHY}"),
      R("[DOPLNIŤ – aktuálny stav ťarchy/výška zostatku úveru ku dňu podpisu]")], space_after=14)

# --- Článok III. Predmet zmluvy ---
heading("Článok III.", 14)
heading("Predmet zmluvy", 12)
body(N(f"Predmetom tejto zmluvy je záväzok Sprostredkovateľa vyvíjať pre {Z_GEN} činnosť smerujúcu k tomu, aby "
       f"{Z_NOM} {MAJU} príležitosť uzatvoriť s tretími osobami zmluvu o budúcej kúpnej zmluve, ktorej súčasťou "
       "bude kúpna zmluva týkajúca sa nehnuteľnosti do dátumu podľa článku V. bodu 1. tejto zmluvy, a to za nižšie "
       f"uvedených podmienok a záväzok {Z_GEN} pri splnení predmetu tejto zmluvy zaplatiť Sprostredkovateľovi Províziu."))
body(N(f"Sprostredkovateľ vykoná nevyhnutnú súčinnosť na zabezpečenie odpredaja nehnuteľnosti tým, že zabezpečí "
       f"potenciálneho kupujúceho, ktorý bude akceptovať podmienky {Z_GEN} a stanovenej budúcej kúpnej ceny."))
body(N(f"Sprostredkovateľ je povinný informovať {Z_GEN} ihneď po obdržaní záujmu a v prípade porušenia exkluzivity "
       f"má voči {Z_DAT} nárok na zmluvnú pokutu vo výške 10 % z výšky Provízie (článok IV., bod 2.)."))
body(N("[POZNÁMKA: Sem doplň text bodov 5.–8. Čl. III zo VZORU — exkluzivita a náhrada nákladov "
       "(vyššie sú len skrátené záväzky).]"))
if DATA["vratane_ponuky_9_12"]:
    body(N("[POZNÁMKA: Sem doplň kompletný text bodov 9.–12. Čl. III zo VZORU — „Ponuka“ "
           "(predám tvoju nehnuteľnosť alebo dovolenka zdarma, 60 dní / dovolenka 1500 EUR). "
           "Používateľ potvrdil, že tieto body MAJÚ byť v zmluve.]"), space_after=14)
else:
    body(N("[POZNÁMKA: Body 9.–12. Čl. III („Ponuka“ — predám tvoju nehnuteľnosť alebo dovolenka zdarma) "
           "boli na žiadosť používateľa VYNECHANÉ pre tohto klienta.]"), space_after=14)

# --- Článok IV. Platobné podmienky ---
heading("Článok IV.", 14)
heading("Platobné podmienky", 12)
body(N(f"1. Cena nehnuteľnosti, za ktorú {MAJU} {Z_NOM} záujem predať nehnuteľnosť, je stanovená na sumu "
       f"{DATA['cena_cislom']} (slovom: {DATA['cena_slovom']})."))
body(N("(ďalej aj ako „cena nehnuteľnosti“)"))
body(N(f"2. Provízia Sprostredkovateľa je stanovená dohodou zmluvných strán na {DATA['provizia_pct']} z celkovej "
       "dohodnutej budúcej kúpnej ceny prevádzanej nehnuteľnosti."))
body(N("(ďalej aj ako „Provízia“)"))
body(N("3. V prípade konania aukcie a/alebo licitovania ceny je dodatočná Provízia 30 % z navýšenia ceny nehnuteľnosti."))
body(N(f"4. {Z_NOM} a Sprostredkovateľ sa dohodli, že suma vo výške Provízie bude prvou časťou budúcej kúpnej ceny, "
       f"ktorú tretia osoba – Budúci kupujúci uhradí pri podpise zmluvy o budúcej kúpnej zmluve na účet {Z_GEN}. "
       f"{Z_NOM} následne túto platbu celú {('poukážu' if PLURAL else 'poukáže')} do 3 dní na účet Sprostredkovateľa "
       "na základe vystavenej faktúry."))
body(N("[POZNÁMKA: Doplň body 5.–8. Čl. IV zo VZORU (spory, vrátenie ceny, provízia po skončení platnosti).]"), space_after=14)

# --- Článok V. Doba trvania ---
heading("Článok V.", 14)
heading("Doba trvania zmluvy", 12)
body(N(f"1. Táto zmluva sa uzatvára na dobu určitú, a to odo dňa jej uzatvorenia na dobu {DATA['doba_dni']} dní."))
body(N("2. Po uplynutí tejto doby môže byť písomnou dohodou všetkých účastníkov predĺžená na ďalšie dohodnuté obdobie."),
     space_after=14)

# --- Článok VI. Spoločné a záverečné ustanovenia ---
heading("Článok VI.", 14)
heading("Spoločné a záverečné ustanovenia", 12)
body(N("Zmluvné strany sa dohodli, že práva a povinnosti neupravené touto zmluvou sa spravujú ust. Občianskeho zákonníka."))
body(N("Meniť alebo doplňovať obsah tejto zmluvy je možné len formou písomných dodatkov podpísaných všetkými zmluvnými stranami."))
body(N(f"Táto zmluva bola vypracovaná v {DATA['pocet_vyhotoveni']} vyhotoveniach, každý Záujemca a Sprostredkovateľ obdržia po jednom vyhotovení."))
body(N("Táto Zmluva nadobúda platnosť a účinnosť dňom podpisu všetkými zmluvnými stranami."))
body(N(f"Bližšie informácie o spracúvaní osobných údajov nájdete na {s['web']} alebo v sídle Sprostredkovateľa."))
body(N("Zmluvné strany prehlasujú, že obsahu zmluvy porozumeli a uzavreli ju podľa svojej slobodnej vôle, čo potvrdzujú svojimi podpismi."),
     space_after=20)

# --- Podpisový blok ---
body([N(f"V {DATA['miesto_podpisu']}, dňa "), R("[DOPLNIŤ DÁTUM PODPISU]")], justify=False, space_after=24)

# Prvý Záujemca + Sprostredkovateľ vedľa seba
table = doc.add_table(rows=4, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
rows = [
    ("Záujemca 1:", "Sprostredkovateľ:"),
    ("", ""),
    ("..............................", ".............................."),
    (ZAUJEMCOVIA[0]["meno"], f"{s['konatel']}, konateľ {s['obchodne_meno']}"),
]
for r, (left, right) in enumerate(rows):
    table.cell(r, 0).text = left
    table.cell(r, 1).text = right
    for c in range(2):
        for p in table.cell(r, c).paragraphs:
            for run in p.runs:
                run.font.name = FONT
                run.font.size = Pt(12)

# Ďalší Záujemcovia (2, 3, ...) každý vlastná čiara
for z in ZAUJEMCOVIA[1:]:
    doc.add_paragraph()
    add_p(f"Záujemca:", space_after=2)
    add_p("..............................", space_after=2)
    add_p(z["meno"], space_after=20)

# --- Príloha: marketingové služby (skopíruj presný text zo VZORU) ---
heading(f"Príklad činností vyvíjaných pre {Z_GEN}:", 13)
body(N(f"VIAC NA: {s['web']}"), justify=False)
body(N("[POZNÁMKA: Sem skopíruj presný zoznam služieb zo VZORU — profesionálny fotograf, video obhliadka, "
       "dron, 3D sken, virtuálny staging, www stránka, katalóg, 3D pôdorys, Facebook, platená reklama, "
       "TikTok, YouTube, aukcie, ďalšie služby.]"))

doc.save(DATA["out_path"])
print("Saved:", DATA["out_path"])
